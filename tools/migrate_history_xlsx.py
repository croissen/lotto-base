"""
lotto.xlsx → Supabase lotto_history (당첨금 컬럼 포함) 마이그레이션

엑셀 컬럼 순서:
  회차, 번호1~6, 보너스, 1등당첨금, 1등당첨자수, 2등당첨금, 2등당첨자수

선행:
  1. sql/05_admin_and_history.sql 실행 (또는 lotto_history 테이블 존재)
  2. sql/06_prize_columns.sql 실행 (당첨금 컬럼 추가됨)
  3. pip install openpyxl psycopg2-binary python-dotenv

실행:
  python tools/migrate_history_xlsx.py [xlsx경로]
  (xlsx경로 생략 시: C:\\Users\\i\\Downloads\\lotto.xlsx)
"""

import os
import sys
import time

try:
    from dotenv import load_dotenv
    import psycopg2
    import openpyxl
except ImportError:
    print("[X] 패키지: pip install openpyxl psycopg2-binary python-dotenv")
    sys.exit(1)


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
ENV_PATH = os.path.join(PROJECT_ROOT, ".env.local")
DEFAULT_XLSX = r"C:\Users\i\Downloads\lotto.xlsx"


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx_path):
        print(f"[X] xlsx 파일 없음: {xlsx_path}")
        sys.exit(1)

    if not os.path.exists(ENV_PATH):
        print(f"[X] .env.local 없음: {ENV_PATH}")
        sys.exit(1)
    load_dotenv(ENV_PATH)
    db_url = os.getenv("SUPABASE_DB_URL")
    if not db_url:
        print("[X] .env.local에 SUPABASE_DB_URL 없음")
        sys.exit(1)

    # 엑셀 로드
    print(f"[+] 엑셀 로드 중: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    print(f"    시트: {ws.title}, 행: {ws.max_row}, 열: {ws.max_column}")

    # 1행은 헤더 — 건너뜀
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        # (회차, n1..n6, bonus, first_amount, first_winners, second_amount, second_winners)
        rows.append((
            int(row[0]),                      # round_no
            int(row[1]), int(row[2]), int(row[3]),
            int(row[4]), int(row[5]), int(row[6]),
            int(row[7]),                      # bonus
            int(row[8])  if row[8]  is not None else None,  # 1등 당첨금
            int(row[9])  if row[9]  is not None else None,  # 1등 당첨자수
            int(row[10]) if row[10] is not None else None,  # 2등 당첨금
            int(row[11]) if row[11] is not None else None,  # 2등 당첨자수
        ))
    print(f"    {len(rows)}개 회차 파싱됨")

    # DB 연결
    print(f"[+] Supabase 연결 중...")
    conn = psycopg2.connect(db_url, connect_timeout=10)
    cur = conn.cursor()

    # 기존 데이터 확인
    cur.execute("SELECT COUNT(*) FROM lotto_history")
    existing = cur.fetchone()[0]
    if existing > 0:
        print(f"[!] lotto_history에 이미 {existing}행 있음.")
        ans = input("    전부 삭제하고 엑셀로 새로 채울까요? (y/n): ").strip().lower()
        if ans != "y":
            print("[-] 취소")
            sys.exit(0)
        cur.execute("TRUNCATE lotto_history")
        conn.commit()
        print(f"    기존 데이터 삭제 완료")

    # INSERT
    print(f"[+] INSERT 중...")
    t0 = time.time()
    cur.executemany(
        """
        INSERT INTO lotto_history (
            round_no, n1, n2, n3, n4, n5, n6, bonus,
            first_prize_amount, first_prize_winners,
            second_prize_amount, second_prize_winners
        ) VALUES (
            %s, %s, %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s
        )
        ON CONFLICT (round_no) DO NOTHING
        """,
        rows,
    )
    conn.commit()
    elapsed = time.time() - t0

    # 결과 확인
    cur.execute("SELECT COUNT(*), MIN(round_no), MAX(round_no) FROM lotto_history")
    cnt, mn, mx = cur.fetchone()
    cur.execute("""
        SELECT round_no, n1, n2, n3, n4, n5, n6, bonus,
               first_prize_amount, first_prize_winners,
               second_prize_amount, second_prize_winners
        FROM lotto_history
        ORDER BY round_no DESC LIMIT 1
    """)
    latest = cur.fetchone()
    cur.close()
    conn.close()

    print(f"\n[+] 마이그레이션 완료!")
    print(f"    소요: {elapsed:.1f}초")
    print(f"    행 수: {cnt}  (회차 {mn} ~ {mx})")
    print(f"\n[샘플] 최신 회차:")
    print(f"    {latest[0]}회: {latest[1:7]} + 보너스 {latest[7]}")
    print(f"    1등 {latest[8]:,}원 × {latest[9]}명")
    print(f"    2등 {latest[10]:,}원 × {latest[11]}명")


if __name__ == "__main__":
    main()
